import shutil
import signal
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from app.services import svn_knowledge as knowledge


def test_legacy_office_files_are_converted_and_indexed_recursively(tmp_path, monkeypatch):
    executable = shutil.which("libreoffice") or shutil.which("soffice")
    if not executable:
        pytest.skip("requires LibreOffice Writer and Calc")
    monkeypatch.setattr(knowledge.settings, "knowledge_root", tmp_path / "knowledge")
    source = tmp_path / "source"
    source.mkdir()
    document = source / "考勤需求.docx"
    with zipfile.ZipFile(document, "w") as archive:
        archive.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>考勤制度：每天记录上下班时间。</w:t></w:r></w:p></w:body></w:document>')
    workbook = Workbook()
    workbook.active.title = "考勤规则"
    workbook.active.append(["迟到次数", 0, False])
    workbook.create_sheet("请假规则").append(["年假天数", 5])
    spreadsheet = source / "请假需求.xlsx"
    workbook.save(spreadsheet)
    workbook.close()
    working_copy = tmp_path / "wc"
    nested = working_copy / "考勤" / "2026 年"
    nested.mkdir(parents=True)
    for path, target in [(document, "doc:MS Word 97"), (spreadsheet, "xls:MS Excel 97")]:
        subprocess.run(
            [executable, "-env:UserInstallation=" + (tmp_path / "generator-profile").as_uri(),
             "--headless", "--convert-to", target, "--outdir", str(nested), str(path)],
            check=True, capture_output=True, timeout=60,
        )
    # Uppercase extensions, Chinese names and spaces must work too.
    (nested / "考勤需求.doc").rename(nested / "考勤需求.DOC")
    originals = {path: path.read_bytes() for path in nested.iterdir()}
    assert len(originals) == 2
    assert all(data.startswith(bytes.fromhex("d0cf11e0a1b11ae1")) for data in originals.values())

    class Embedding:
        base_url = "http://embedding.example/v1"
        model = "test"

        def embed(self, texts):
            return [[1.0, 0.0] for _ in texts]

    copies = {"部门制度": working_copy}
    manifest, changes = knowledge._build_manifest("https://svn.example/repo", {"部门制度": "1"}, copies, {})
    assert changes["added"] == 2
    knowledge._publish_vector_index(manifest, {}, copies, Embedding())
    assert manifest["failed_files"] == {}
    results = knowledge.search_vector_index("考勤", [1.0, 0.0], 10, full_content=True)
    content = {item["source_path"]: item["content"] for item in results}
    assert "每天记录上下班时间" in content["部门制度/考勤/2026 年/考勤需求.DOC"]
    sheet_text = content["部门制度/考勤/2026 年/请假需求.xls"]
    assert "工作表：考勤规则" in sheet_text and "迟到次数\t0\tFalse" in sheet_text
    assert "工作表：请假规则" in sheet_text and "年假天数\t5" in sheet_text
    assert all(path.read_bytes() == data for path, data in originals.items())
    _, changes = knowledge._build_manifest("https://svn.example/repo", {"部门制度": "1"}, copies, manifest)
    assert changes == {"added": 0, "changed": 0, "deleted": 0, "unchanged": 2}


@pytest.mark.parametrize("suffix", [".doc", ".xls"])
def test_legacy_office_missing_dependency_is_reported(tmp_path, monkeypatch, suffix):
    path = tmp_path / ("需求" + suffix)
    path.write_bytes(b"legacy")
    monkeypatch.setattr(knowledge.shutil, "which", lambda name: None)
    with pytest.raises(ValueError, match="需要 LibreOffice"):
        knowledge._extract_text(path)


@pytest.mark.parametrize("failure", ["timeout", "exit", "no_output"])
def test_legacy_conversion_failure_cleans_up_and_reports_error(tmp_path, monkeypatch, failure):
    path = tmp_path / "需求.doc"
    path.write_bytes(b"legacy")
    monkeypatch.setattr(knowledge.shutil, "which", lambda name: "/usr/bin/libreoffice")
    killed = []
    monkeypatch.setattr(knowledge.os, "killpg", lambda pid, sig: killed.append((pid, sig)))
    output_roots = []

    class Process:
        pid = 123
        returncode = 1 if failure == "exit" else 0

        def __init__(self, command, **kwargs):
            root = Path(command[command.index("--outdir") + 1])
            output_roots.append(root)
            assert (root / "profile" / "registrymodifications.xcu").is_file()
            assert kwargs["start_new_session"] is True

        def __enter__(self):
            return self

        def __exit__(self, *args):
            pass

        def wait(self, timeout=None):
            if failure == "timeout" and timeout is not None:
                raise subprocess.TimeoutExpired("libreoffice", timeout)

    monkeypatch.setattr(knowledge.subprocess, "Popen", Process)
    with pytest.raises(ValueError, match="转换超时" if failure == "timeout" else "转换失败"):
        knowledge._extract_text(path)
    assert killed == ([(123, signal.SIGKILL)] if failure == "timeout" else [])
    assert all(not root.exists() for root in output_roots)
    assert path.read_bytes() == b"legacy"


def test_office_size_and_text_boundaries(tmp_path, monkeypatch):
    from types import SimpleNamespace
    # Real parsers with simulated file sizes avoid allocating 500 MiB in the test.
    doc = tmp_path / 'large.DOCX'
    with zipfile.ZipFile(doc, 'w') as archive:
        archive.writestr('word/document.xml', '<document><text>正文</text></document>')
    sheet = tmp_path / 'large.xlsx'
    workbook = Workbook()
    workbook.active.append(['正文'])
    workbook.save(sheet)
    workbook.close()
    original_stat = Path.stat
    size = 500 * 1024 * 1024
    monkeypatch.setattr(Path, 'stat', lambda path, *args, **kwargs: SimpleNamespace(st_size=size) if path in (doc, sheet) else original_stat(path, *args, **kwargs))
    for path in (doc, sheet):
        assert '正文' in knowledge._extract_text(path)
    size += 1
    for path in (doc, sheet):
        with pytest.raises(ValueError, match='500 MiB'):
            knowledge._extract_text(path)
    assert knowledge.file_size_limit('other.pdf') == 50 * 1024 * 1024
    assert sum(map(len, knowledge._chunks('文' * 10_000_000, size=100_000, overlap=0))) == 10_000_000
    with pytest.raises(ValueError, match='1000 万'):
        list(knowledge._chunks('文' * 10_000_001))
